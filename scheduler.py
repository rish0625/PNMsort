#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FRATERNITY INTERVIEW SCHEDULER
==============================

Randomly groups PNMs ("rushes") into groups of exactly 4 for 20-minute group
interviews, using only the availability they submitted through a Google Form.

This program knows NOTHING about brothers/interviewers. Its only inputs are
PNM names, PNM availability, the list of interview slots, and scheduling
constraints. The five brothers per group are assigned by hand afterward, in
the blank Brother columns of the Excel output.

Run it with:

    python scheduler.py

Requires: openpyxl  (pip install -r requirements.txt)
"""

import argparse
import csv
import os
import random
import re
import sys
from datetime import datetime, timedelta

sys.setrecursionlimit(10000)

# =============================================================================
# ============================ CONFIGURATION ==================================
# =============================================================================
# EDIT THIS SECTION. Everything you are likely to change lives here.
# Anything you enter at the prompts when running the program overrides these.

# --- Group and timing rules -------------------------------------------------
GROUP_SIZE = 4            # PNMs per interview group. Exactly this many, always.
INTERVIEW_DURATION = 20   # Minutes per interview.
MAX_GROUPS_PER_SLOT = 1   # How many groups may interview at the same time
                          # (i.e. how many interview rooms you have).
MAX_ATTEMPTS = 10000      # How many random shuffles to try before giving up.

# --- Which CSV columns to read ----------------------------------------------
# These must match the column headers in the CSV exported from Google Forms.
# Matching is case-insensitive and ignores extra spaces.
NAME_COLUMN = "Name"
AVAILABILITY_COLUMN = "Interview Availability"

# --- The interview slots ----------------------------------------------------
# One entry per 20-minute interview slot, listed in chronological order.
# These strings must match the checkbox options in your Google Form.
# Each entry is the START of a slot; the end time is computed automatically
# using INTERVIEW_DURATION.
INTERVIEW_SLOTS = [
    "Monday 6:00 PM",
    "Monday 6:20 PM",
    "Monday 6:40 PM",
    "Monday 7:00 PM",
    "Monday 7:20 PM",
    "Monday 7:40 PM",
    "Tuesday 6:00 PM",
    "Tuesday 6:20 PM",
    "Tuesday 6:40 PM",
    "Tuesday 7:00 PM",
    "Tuesday 7:20 PM",
    "Tuesday 7:40 PM",
]

# --- Default filenames shown at the prompts ---------------------------------
DEFAULT_CSV = "PNM_Availability.csv"
DEFAULT_OUTPUT = "Interview_Schedule.xlsx"

# =============================================================================
# ======================= END OF CONFIGURATION ================================
# =============================================================================


class SchedulerError(Exception):
    """A problem the user can understand and fix. Printed without a traceback."""


# -----------------------------------------------------------------------------
# Time slot handling
# -----------------------------------------------------------------------------

TIME_PATTERN = re.compile(
    r"^(?P<prefix>.*?)\s*(?P<hour>\d{1,2})\s*[:.]\s*(?P<minute>\d{2})\s*"
    r"(?P<meridiem>[AaPp]\.?\s*[Mm]\.?)?\s*$"
)


def _format_time(dt):
    hour = dt.hour % 12 or 12
    meridiem = "AM" if dt.hour < 12 else "PM"
    return "{}:{:02d} {}".format(hour, dt.minute, meridiem)


def normalize_text(text):
    """Upper-case, collapse whitespace, drop stray punctuation differences."""
    text = (text or "").replace("\u2013", "-").replace("\u2014", "-")
    text = text.replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip().upper()


def strip_range(token):
    """'Monday 6:00 PM - 6:20 PM' -> 'Monday 6:00 PM'. Also handles ' to '."""
    token = token.replace("\u2013", "-").replace("\u2014", "-")
    token = re.split(r"\s+-\s+|\s+to\s+", token, flags=re.IGNORECASE)[0]
    return token.strip()


class Slot(object):
    """One 20-minute interview slot."""

    def __init__(self, index, label, duration):
        self.index = index
        self.label = label.strip()

        match = TIME_PATTERN.match(self.label)
        if not match:
            raise SchedulerError(
                "The interview slot '{}' in INTERVIEW_SLOTS could not be read.\n"
                "Each slot needs a clock time, for example:\n"
                "    \"Monday 6:00 PM\"   or   \"Monday 18:00\"".format(self.label)
            )

        hour = int(match.group("hour"))
        minute = int(match.group("minute"))
        meridiem = match.group("meridiem")

        if meridiem:
            meridiem = meridiem.replace(".", "").replace(" ", "").upper()
            if not 1 <= hour <= 12:
                raise SchedulerError(
                    "The interview slot '{}' has an hour of {}, which is not valid "
                    "with AM/PM.".format(self.label, hour)
                )
            if meridiem == "PM" and hour != 12:
                hour += 12
            elif meridiem == "AM" and hour == 12:
                hour = 0
        elif not 0 <= hour <= 23:
            raise SchedulerError(
                "The interview slot '{}' has an invalid hour.".format(self.label)
            )

        if not 0 <= minute <= 59:
            raise SchedulerError(
                "The interview slot '{}' has an invalid minute.".format(self.label)
            )

        self.day = match.group("prefix").strip().rstrip(",").strip()
        self.start = datetime(2000, 1, 1, hour, minute)
        self.end = self.start + timedelta(minutes=duration)

    @property
    def start_label(self):
        return ("{} {}".format(self.day, _format_time(self.start))).strip()

    @property
    def end_label(self):
        return ("{} {}".format(self.day, _format_time(self.end))).strip()

    @property
    def window_label(self):
        return "{} - {}".format(self.start_label, _format_time(self.end))

    def __repr__(self):
        return "<Slot {}>".format(self.window_label)


def build_slots(slot_labels, duration):
    if not slot_labels:
        raise SchedulerError(
            "No interview slots are configured.\n"
            "Open scheduler.py and add slots to the INTERVIEW_SLOTS list."
        )

    slots = []
    seen = {}
    for index, label in enumerate(slot_labels):
        slot = Slot(index, label, duration)
        key = normalize_text(slot.label)
        if key in seen:
            raise SchedulerError(
                "The interview slot '{}' is listed twice in INTERVIEW_SLOTS.\n"
                "Remove the duplicate and run the program again.".format(slot.label)
            )
        seen[key] = True
        slots.append(slot)
    return slots


def build_slot_lookup(slots):
    """Map several spellings of a slot to its index, for matching form answers."""
    lookup = {}
    for slot in slots:
        for variant in (slot.label, slot.start_label, slot.window_label):
            lookup[normalize_text(variant)] = slot.index
        # Also allow a bare time, but only if it is unambiguous across days.
        bare = normalize_text(_format_time(slot.start))
        lookup.setdefault(bare, slot.index)
        if lookup[bare] != slot.index:
            lookup[bare] = None  # ambiguous; refuse to guess
    return {k: v for k, v in lookup.items() if v is not None}


# -----------------------------------------------------------------------------
# Reading the CSV
# -----------------------------------------------------------------------------

class PNM(object):
    def __init__(self, name, raw_availability, slot_indexes):
        self.name = name
        self.raw_availability = raw_availability
        self.slot_indexes = frozenset(slot_indexes)


def find_column(fieldnames, wanted, description):
    normalized = {normalize_text(f): f for f in fieldnames if f is not None}
    target = normalize_text(wanted)

    if target in normalized:
        return normalized[target]

    partial = [orig for norm, orig in normalized.items() if target in norm]
    if len(partial) == 1:
        return partial[0]

    columns = "\n".join("    - {}".format(f) for f in fieldnames if f)
    raise SchedulerError(
        "Could not find the {} column.\n\n"
        "The program is looking for a column named:\n"
        "    {}\n\n"
        "The CSV file actually contains these columns:\n{}\n\n"
        "Fix this by either renaming the column in the CSV, or by editing the\n"
        "CONFIGURATION section at the top of scheduler.py.".format(
            description, wanted, columns or "    (no columns found)"
        )
    )


def split_availability(cell):
    """Split a Google Forms checkbox cell into individual slot answers."""
    cell = (cell or "").replace("\n", ",").replace(";", ",")
    return [part.strip() for part in cell.split(",") if part.strip()]


def match_availability(tokens, slot_lookup):
    """Turn raw answers into slot indexes. Returns (indexes, unrecognized)."""
    indexes = []
    unrecognized = []

    i = 0
    while i < len(tokens):
        token = tokens[i]
        index = slot_lookup.get(normalize_text(strip_range(token)))

        # Recovery: a form option containing a comma ("Monday, 6:00 PM") gets
        # split in half, so try re-joining this token with the next one.
        if index is None and i + 1 < len(tokens):
            joined = "{}, {}".format(token, tokens[i + 1])
            joined_index = slot_lookup.get(normalize_text(strip_range(joined)))
            if joined_index is not None:
                indexes.append(joined_index)
                i += 2
                continue

        if index is None:
            unrecognized.append(token)
        else:
            indexes.append(index)
        i += 1

    return sorted(set(indexes)), unrecognized


def load_pnms(csv_path, name_column, availability_column, slots):
    if not os.path.isfile(csv_path):
        raise SchedulerError(
            "Could not find the CSV file:\n    {}\n\n"
            "Check that the file name is spelled correctly and that the file is\n"
            "in the same folder as scheduler.py (or give the full path to it).".format(
                os.path.abspath(csv_path)
            )
        )

    try:
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            fieldnames = reader.fieldnames
            rows = list(reader)
    except UnicodeDecodeError:
        with open(csv_path, "r", encoding="latin-1", newline="") as handle:
            reader = csv.DictReader(handle)
            fieldnames = reader.fieldnames
            rows = list(reader)
    except OSError as exc:
        raise SchedulerError("Could not open the CSV file:\n    {}".format(exc))

    if not fieldnames:
        raise SchedulerError(
            "The CSV file appears to be empty or is not a valid CSV file:\n    {}\n\n"
            "Export the responses from Google Forms again "
            "(File > Download > Comma Separated Values).".format(csv_path)
        )

    name_field = find_column(fieldnames, name_column, "PNM name")
    availability_field = find_column(
        fieldnames, availability_column, "availability"
    )

    if not rows:
        raise SchedulerError(
            "The CSV file has column headers but no responses in it:\n    {}".format(
                csv_path
            )
        )

    slot_lookup = build_slot_lookup(slots)

    pnms = []
    empty_name_rows = []
    no_availability = []
    bad_values = []
    seen_names = {}
    duplicates = []

    for row_number, row in enumerate(rows, start=2):  # row 1 is the header
        name = (row.get(name_field) or "").strip()
        name = re.sub(r"\s+", " ", name)
        raw = (row.get(availability_field) or "").strip()

        if not name:
            if raw:
                empty_name_rows.append(row_number)
            continue  # a completely blank row is just ignored

        key = name.upper()
        if key in seen_names:
            duplicates.append((name, seen_names[key], row_number))
            continue
        seen_names[key] = row_number

        tokens = split_availability(raw)
        indexes, unrecognized = match_availability(tokens, slot_lookup)

        for value in unrecognized:
            bad_values.append((name, value))
        if not indexes and not unrecognized:
            no_availability.append(name)

        pnms.append(PNM(name, raw, indexes))

    problems = []

    if empty_name_rows:
        problems.append(
            "These rows have availability but no name (row numbers as shown in "
            "Excel):\n    {}\n  Add the missing names to the CSV.".format(
                ", ".join(str(r) for r in empty_name_rows)
            )
        )

    if duplicates:
        lines = [
            "    {}  (rows {} and {})".format(name, first, second)
            for name, first, second in duplicates
        ]
        problems.append(
            "These PNM names appear more than once:\n{}\n"
            "  Each PNM must appear exactly once. If two different people share a\n"
            "  name, make them distinct (for example 'John Smith' and "
            "'John Smith Jr.').".format("\n".join(lines))
        )

    if bad_values:
        lines = []
        for name, value in bad_values[:20]:
            lines.append("    {}: \"{}\"".format(name, value))
        if len(bad_values) > 20:
            lines.append("    ...and {} more".format(len(bad_values) - 20))
        valid = "\n".join("    - {}".format(s.label) for s in slots)
        problems.append(
            "These availability answers do not match any configured interview "
            "slot:\n{}\n\n  The configured slots are:\n{}\n\n"
            "  Fix this by editing INTERVIEW_SLOTS at the top of scheduler.py so "
            "the\n  slots match the checkbox options in your Google Form "
            "exactly.".format("\n".join(lines), valid)
        )

    if no_availability:
        problems.append(
            "These PNMs did not select any availability:\n{}\n"
            "  Ask them to resubmit the form, or remove them from the CSV before "
            "running\n  the scheduler again.".format(
                "\n".join("    - {}".format(n) for n in no_availability)
            )
        )

    if problems:
        raise SchedulerError(
            "The CSV file has {} problem(s) that must be fixed:\n\n{}".format(
                len(problems),
                "\n\n".join(
                    "{}. {}".format(i, text) for i, text in enumerate(problems, 1)
                ),
            )
        )

    if not pnms:
        raise SchedulerError(
            "No PNMs were found in the CSV file:\n    {}\n\n"
            "Check that the '{}' column actually contains names.".format(
                csv_path, name_field
            )
        )

    return pnms


# -----------------------------------------------------------------------------
# Randomized grouping and scheduling
# -----------------------------------------------------------------------------

class Group(object):
    def __init__(self, group_id, members, slot=None):
        self.group_id = group_id
        self.members = members
        self.slot = slot

    @property
    def common_slots(self):
        common = None
        for member in self.members:
            common = (
                set(member.slot_indexes)
                if common is None
                else common & member.slot_indexes
            )
        return common or set()


def make_group_ids(count, rng):
    """Neutral IDs like G-7F3A. They encode nothing about the PNMs."""
    ids = set()
    while len(ids) < count:
        ids.add("G-{:04X}".format(rng.randrange(0x10000)))
    return sorted(ids)


def partition(pnms, group_size):
    return [
        pnms[i:i + group_size] for i in range(0, len(pnms), group_size)
    ]


def _augment(group_index, adjacency, slot_node_owner, visited):
    """Standard augmenting-path search (Kuhn's algorithm)."""
    for node in adjacency[group_index]:
        if node in visited:
            continue
        visited.add(node)
        owner = slot_node_owner.get(node)
        if owner is None or _augment(owner, adjacency, slot_node_owner, visited):
            slot_node_owner[node] = group_index
            return True
    return False


def assign_slots(member_groups, slots, capacity, rng):
    """
    Try to give every group a slot that all of its members are available for,
    without exceeding `capacity` groups in any one slot.

    Returns a list of slot indexes (one per group) or None if impossible.
    Each slot is modelled as `capacity` interchangeable seats, so this is an
    ordinary bipartite matching problem, solved exactly rather than greedily.
    """
    adjacency = []
    for members in member_groups:
        common = None
        for member in members:
            common = (
                set(member.slot_indexes)
                if common is None
                else common & member.slot_indexes
            )
        common = sorted(common or set())
        if not common:
            return None
        nodes = [(slot_index, seat)
                 for slot_index in common
                 for seat in range(capacity)]
        rng.shuffle(nodes)          # no deterministic preference for early slots
        adjacency.append(nodes)

    order = list(range(len(member_groups)))
    rng.shuffle(order)              # no deterministic preference for early groups

    slot_node_owner = {}
    for group_index in order:
        if not _augment(group_index, adjacency, slot_node_owner, set()):
            return None

    assignment = [None] * len(member_groups)
    for (slot_index, _seat), group_index in slot_node_owner.items():
        assignment[group_index] = slot_index
    return assignment


def constrained_attempt(pnms, slots, group_size, capacity, rng):
    """
    Second search strategy, used only when plain reshuffling keeps failing.

    Instead of shuffling PNMs and hoping the groups happen to line up, this
    picks a random set of interview slots to run, then randomly deals every PNM
    into a seat at a slot they said they were free for. It uses exactly the same
    objective constraints as the shuffle method; it just draws its random
    grouping from the valid ones instead of from all possible ones.

    Returns a list of (slot, members) or None if this random draw failed.
    """
    group_count = len(pnms) // group_size

    # Randomly choose which slots will actually host a group.
    instances = [
        (slot.index, seat) for slot in slots for seat in range(capacity)
    ]
    rng.shuffle(instances)
    hosted = [slot_index for slot_index, _seat in instances[:group_count]]

    # Deal PNMs into seats: group g has `group_size` interchangeable seats.
    adjacency = []
    for pnm in pnms:
        seats = [
            (g, seat)
            for g in range(group_count) if hosted[g] in pnm.slot_indexes
            for seat in range(group_size)
        ]
        if not seats:
            return None
        rng.shuffle(seats)
        adjacency.append(seats)

    order = list(range(len(pnms)))
    rng.shuffle(order)

    seat_owner = {}
    for pnm_index in order:
        if not _augment(pnm_index, adjacency, seat_owner, set()):
            return None

    members = [[] for _ in range(group_count)]
    for (group_index, _seat), pnm_index in seat_owner.items():
        members[group_index].append(pnms[pnm_index])

    # Groups sharing a slot are interchangeable, so reshuffle across them to
    # avoid any bias left over from the order the seats were filled in.
    by_slot = {}
    for group_index, slot_index in enumerate(hosted):
        by_slot.setdefault(slot_index, []).append(group_index)
    for slot_index, group_indexes in by_slot.items():
        if len(group_indexes) > 1:
            pooled = [m for g in group_indexes for m in members[g]]
            rng.shuffle(pooled)
            for position, group_index in enumerate(group_indexes):
                members[group_index] = pooled[
                    position * group_size:(position + 1) * group_size
                ]

    return [(slots[hosted[g]], members[g]) for g in range(group_count)]


def generate_schedule(pnms, slots, group_size, capacity, max_attempts, rng):
    """
    Search for a valid randomized schedule.

    Phase 1 (plain reshuffling): shuffle every PNM, cut the deck into groups of
    four, and see whether every group shares a free slot. This ignores
    availability entirely while forming groups, so it is the most obviously
    impartial method, and it is what runs in the normal case.

    Phase 2 (availability-aware random dealing): if thousands of blind shuffles
    all fail, availability is too tight for luck alone. Phase 2 then draws a
    random grouping from among the ones that actually work. Same constraints,
    same absence of preferences about who is with whom - it just stops throwing
    away draws that were never going to schedule.
    """
    pool = sorted(pnms, key=lambda p: p.name.upper())  # stable starting order
    attempts = 0

    def finish(pairs, attempts_used, method):
        group_ids = make_group_ids(len(pairs), rng)
        rng.shuffle(group_ids)
        groups = [
            Group(group_ids[i], members, slot)
            for i, (slot, members) in enumerate(pairs)
        ]
        groups.sort(key=lambda g: (g.slot.index, g.group_id))
        return groups, attempts_used, method

    for attempts in range(1, max_attempts + 1):
        shuffled = pool[:]
        rng.shuffle(shuffled)
        member_groups = partition(shuffled, group_size)
        assignment = assign_slots(member_groups, slots, capacity, rng)
        if assignment is not None:
            pairs = [
                (slots[assignment[i]], member_groups[i])
                for i in range(len(member_groups))
            ]
            return finish(pairs, attempts, "Random shuffle")

    phase_two_budget = max(200, max_attempts // 10)
    for extra in range(1, phase_two_budget + 1):
        shuffled = pool[:]
        rng.shuffle(shuffled)
        pairs = constrained_attempt(shuffled, slots, group_size, capacity, rng)
        if pairs is not None:
            return finish(
                pairs, attempts + extra,
                "Availability-constrained random draw "
                "(after {:,} blind shuffles failed)".format(attempts),
            )
        attempts_total = attempts + extra

    raise SchedulerError(build_failure_message(
        pnms, slots, group_size, capacity, attempts + phase_two_budget
    ))


def build_failure_message(pnms, slots, group_size, capacity, attempts):
    """Explain a scheduling failure in terms the user can act on."""
    coverage = []
    for slot in slots:
        count = sum(1 for p in pnms if slot.index in p.slot_indexes)
        coverage.append((slot, count))

    tightest = sorted(pnms, key=lambda p: (len(p.slot_indexes), p.name.upper()))[:8]
    tight_lines = "\n".join(
        "    - {} ({} slot{})".format(
            p.name, len(p.slot_indexes), "" if len(p.slot_indexes) == 1 else "s"
        )
        for p in tightest
    )
    thin = [
        "    - {}: only {} PNM(s) available".format(slot.window_label, count)
        for slot, count in coverage
        if count < group_size
    ]

    message = (
        "Tried {:,} random groupings - both blind shuffles and "
        "availability-aware\ndraws - and could not find one where every group of "
        "{} shares an interview\nslot.\n\n"
        "This usually means the availability is too tight, not that the program "
        "failed.\n\n"
        "PNMs who submitted the fewest available slots:\n{}\n".format(
            attempts, group_size, tight_lines
        )
    )
    if thin:
        message += (
            "\nSlots that fewer than {} PNMs selected (these can never host a "
            "group):\n{}\n".format(group_size, "\n".join(thin))
        )
    message += (
        "\nThings that will fix this:\n"
        "  1. Add more interview slots to INTERVIEW_SLOTS in scheduler.py, and ask\n"
        "     PNMs to resubmit availability for them.\n"
        "  2. Increase 'Maximum groups per slot' if you have more than one "
        "interview room.\n"
        "  3. Follow up with the PNMs listed above and ask them to open up more "
        "times.\n"
        "  4. Raise MAX_ATTEMPTS if you believe a valid schedule exists but is "
        "rare."
    )
    return message


def validate_inputs(pnms, slots, group_size, capacity):
    count = len(pnms)

    if group_size < 1:
        raise SchedulerError("Group size must be at least 1.")

    if capacity < 1:
        raise SchedulerError("Maximum groups per slot must be at least 1.")

    if count % group_size != 0:
        remainder = count % group_size
        lower = count - remainder
        upper = lower + group_size
        raise SchedulerError(
            "There are {} PNMs, which cannot be split into groups of exactly {}.\n\n"
            "{} divided by {} leaves {} PNM(s) left over, and this program will not\n"
            "silently create a group of {} or {}.\n\n"
            "Workable numbers near {} are {} PNMs ({} groups) and {} PNMs ({} "
            "groups).\n\n"
            "Fix this by adding or removing PNMs in the CSV, or by changing the "
            "group\nsize when the program asks for it.".format(
                count, group_size, count, group_size, remainder,
                group_size - 1, group_size + 1,
                count, lower, lower // group_size, upper, upper // group_size,
            )
        )

    groups_needed = count // group_size
    seats = len(slots) * capacity
    if groups_needed > seats:
        raise SchedulerError(
            "There is not enough interview capacity.\n\n"
            "    PNMs:                    {}\n"
            "    Groups needed:           {}\n"
            "    Interview slots:         {}\n"
            "    Groups allowed per slot: {}\n"
            "    Total capacity:          {} group(s)\n\n"
            "Fix this by adding more slots to INTERVIEW_SLOTS at the top of\n"
            "scheduler.py, or by raising 'Maximum groups per slot' if you have "
            "more\nthan one interview room available.".format(
                count, groups_needed, len(slots), capacity, seats
            )
        )

    return groups_needed


def verify_schedule(groups, pnms, group_size, capacity):
    """A final self-check, so a bad schedule can never reach the Excel file."""
    assigned = []
    per_slot = {}

    for group in groups:
        if len(group.members) != group_size:
            raise SchedulerError(
                "Internal check failed: group {} has {} PNMs instead of {}.".format(
                    group.group_id, len(group.members), group_size
                )
            )
        for member in group.members:
            assigned.append(member.name)
            if group.slot.index not in member.slot_indexes:
                raise SchedulerError(
                    "Internal check failed: {} was assigned to {}, which they did "
                    "not select.".format(member.name, group.slot.window_label)
                )
        per_slot[group.slot.index] = per_slot.get(group.slot.index, 0) + 1

    for slot_index, count in per_slot.items():
        if count > capacity:
            raise SchedulerError(
                "Internal check failed: {} groups were placed in one slot, but the "
                "limit is {}.".format(count, capacity)
            )

    if len(assigned) != len(set(n.upper() for n in assigned)):
        raise SchedulerError("Internal check failed: a PNM was assigned twice.")
    if len(assigned) != len(pnms):
        raise SchedulerError(
            "Internal check failed: {} of {} PNMs were assigned.".format(
                len(assigned), len(pnms)
            )
        )


# -----------------------------------------------------------------------------
# Excel output
# -----------------------------------------------------------------------------

def write_excel(path, groups, pnms, slots, info):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise SchedulerError(
            "The 'openpyxl' library is not installed, so the Excel file cannot be\n"
            "created.\n\nInstall it by running:\n\n    pip install openpyxl\n\n"
            "or:\n\n    pip install -r requirements.txt"
        )

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F3864")
    brother_fill = PatternFill("solid", fgColor="FFF2CC")
    stripe_fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    workbook = Workbook()

    def style_header(sheet, columns):
        for column_index, (title, width) in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=column_index, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = centered
            cell.border = border
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        sheet.row_dimensions[1].height = 28
        sheet.freeze_panes = "A2"

    # ---------------- Sheet 1: Interview Schedule ----------------
    schedule = workbook.active
    schedule.title = "Interview Schedule"

    columns = [
        ("Interview Date/Time", 22), ("End Time", 22), ("Group ID", 11),
    ]
    # One PNM column per group member (four by default).
    columns += [("PNM {}".format(i + 1), 20) for i in range(info["group_size"])]
    columns += [("Brother {}".format(i), 18) for i in range(1, 6)]
    style_header(schedule, columns)

    first_brother_column = 4 + info["group_size"]

    for row_index, group in enumerate(groups, start=2):
        names = sorted(m.name for m in group.members)
        values = [group.slot.start_label, group.slot.end_label, group.group_id]
        values += names
        values += [""] * 5

        for column_index, value in enumerate(values, start=1):
            cell = schedule.cell(row=row_index, column=column_index, value=value)
            cell.border = border
            cell.alignment = centered if column_index <= 3 else left
            if column_index >= first_brother_column:
                cell.fill = brother_fill
            elif row_index % 2 == 0:
                cell.fill = stripe_fill
        schedule.row_dimensions[row_index].height = 20

    schedule.print_title_rows = "1:1"
    schedule.page_setup.orientation = "landscape"
    schedule.page_setup.fitToWidth = 1
    schedule.page_setup.fitToHeight = 0
    schedule.sheet_properties.pageSetUpPr.fitToPage = True

    note_row = len(groups) + 3
    note = schedule.cell(
        row=note_row, column=1,
        value="The Brother columns are intentionally blank. Assign five brothers "
              "to each group by hand.",
    )
    note.font = Font(italic=True, color="7F7F7F")

    # ---------------- Sheet 2: PNM Details ----------------
    details = workbook.create_sheet("PNM Details")
    detail_columns = [
        ("PNM Name", 24), ("Group ID", 11), ("Interview Start", 20),
        ("Interview End", 20), ("Available Times", 70),
    ]
    style_header(details, detail_columns)

    rows = []
    for group in groups:
        for member in group.members:
            available = ", ".join(
                slots[i].label for i in sorted(member.slot_indexes)
            )
            rows.append([
                member.name, group.group_id, group.slot.start_label,
                group.slot.end_label, available,
            ])
    rows.sort(key=lambda r: r[0].upper())

    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            cell = details.cell(row=row_index, column=column_index, value=value)
            cell.border = border
            cell.alignment = left if column_index in (1, 5) else centered
            if row_index % 2 == 0:
                cell.fill = stripe_fill

    details.print_title_rows = "1:1"
    details.page_setup.orientation = "landscape"
    details.page_setup.fitToWidth = 1
    details.page_setup.fitToHeight = 0
    details.sheet_properties.pageSetUpPr.fitToPage = True

    # ---------------- Sheet 3: Randomization Info ----------------
    audit = workbook.create_sheet("Randomization Info")
    style_header(audit, [("Item", 30), ("Value", 46)])

    entries = [
        ("Randomization Seed", info["seed"]),
        ("Date/Time Generated", info["generated"]),
        ("Number of PNMs", info["pnm_count"]),
        ("Group Size", info["group_size"]),
        ("Number of Groups", info["group_count"]),
        ("Interview Duration (minutes)", info["duration"]),
        ("Number of Interview Slots", info["slot_count"]),
        ("Maximum Groups Per Slot", info["capacity"]),
        ("Number of Attempts", info["attempts"]),
        ("Grouping Method", info["method"]),
        ("Source CSV File", info["csv_path"]),
    ]
    for row_index, (label, value) in enumerate(entries, start=2):
        label_cell = audit.cell(row=row_index, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.border = border
        label_cell.alignment = left
        value_cell = audit.cell(row=row_index, column=2, value=value)
        value_cell.border = border
        value_cell.alignment = left

    reproduce = audit.cell(
        row=len(entries) + 3, column=1,
        value="To reproduce this exact schedule, run scheduler.py again with the "
              "same CSV file and enter the seed above when prompted.",
    )
    reproduce.font = Font(italic=True, color="7F7F7F")

    try:
        workbook.save(path)
    except OSError as exc:
        raise SchedulerError(
            "Could not save the Excel file:\n    {}\n\n{}\n\n"
            "If the file is currently open in Excel, close it and run the program "
            "again.".format(os.path.abspath(path), exc)
        )


# -----------------------------------------------------------------------------
# Command line / prompts
# -----------------------------------------------------------------------------

def ask(prompt, default=None, note=None):
    suffix = " [{}]".format(default) if default not in (None, "") else ""
    print("\n{}{}:".format(prompt, suffix))
    if note:
        print("  ({})".format(note))
    try:
        answer = input("> ").strip()
    except EOFError:
        answer = ""
    if not answer and default is not None:
        return str(default)
    return answer


def ask_int(prompt, default, minimum=1, note=None):
    while True:
        answer = ask(prompt, default, note)
        try:
            value = int(answer)
        except ValueError:
            print("  That is not a whole number. Please try again.")
            continue
        if value < minimum:
            print("  Please enter a number of at least {}.".format(minimum))
            continue
        return value


def parse_args():
    parser = argparse.ArgumentParser(
        description="Randomly group PNMs into interview groups of 4.",
        epilog="Run with no options to be prompted for everything.",
    )
    parser.add_argument("--csv", help="Input CSV exported from Google Forms")
    parser.add_argument("--output", help="Output .xlsx filename")
    parser.add_argument("--seed", type=int, help="Randomization seed to reuse")
    parser.add_argument("--group-size", type=int, help="PNMs per group")
    parser.add_argument("--duration", type=int, help="Interview length in minutes")
    parser.add_argument("--capacity", type=int, help="Max groups per slot")
    parser.add_argument("--max-attempts", type=int, help="Max random attempts")
    return parser.parse_args()


def run(args):
    print("=" * 40)
    print("FRATERNITY INTERVIEW SCHEDULER")
    print("=" * 40)

    interactive = args.csv is None

    csv_path = args.csv
    if csv_path is None:
        while True:
            csv_path = ask("Enter CSV filename", DEFAULT_CSV)
            if os.path.isfile(csv_path):
                break
            print("\n  There is no file named '{}' in this folder.".format(csv_path))
            print("  Check the spelling, or paste the full path to the file.")

    output_path = args.output or (
        ask("Enter output filename", DEFAULT_OUTPUT) if interactive
        else DEFAULT_OUTPUT
    )
    if not output_path.lower().endswith(".xlsx"):
        output_path += ".xlsx"

    group_size = args.group_size or (
        ask_int("Group size", GROUP_SIZE) if interactive else GROUP_SIZE
    )
    duration = args.duration or (
        ask_int("Interview duration (minutes)", INTERVIEW_DURATION)
        if interactive else INTERVIEW_DURATION
    )
    capacity = args.capacity or (
        ask_int(
            "Maximum groups per {}-minute slot".format(duration),
            MAX_GROUPS_PER_SLOT,
            note="set this to the number of interview rooms you have",
        ) if interactive else MAX_GROUPS_PER_SLOT
    )
    max_attempts = args.max_attempts or MAX_ATTEMPTS

    seed = args.seed
    if seed is None and interactive:
        while True:
            answer = ask(
                "Randomization seed", "",
                note="leave blank for a new random seed",
            )
            if not answer:
                break
            try:
                seed = int(answer)
                break
            except ValueError:
                print("  The seed must be a whole number, such as 847293.")
    if seed is None:
        seed = random.SystemRandom().randrange(100000, 1000000)

    rng = random.Random(seed)

    print("\nProcessing...\n")

    slots = build_slots(INTERVIEW_SLOTS, duration)
    pnms = load_pnms(csv_path, NAME_COLUMN, AVAILABILITY_COLUMN, slots)
    groups_needed = validate_inputs(pnms, slots, group_size, capacity)

    print("{} PNMs found.".format(len(pnms)))
    print("{} groups required.".format(groups_needed))
    print("{} interview slots of {} minutes are configured.".format(
        len(slots), duration))

    print("\nGenerating randomized schedule...")
    groups, attempts, method = generate_schedule(
        pnms, slots, group_size, capacity, max_attempts, rng
    )
    verify_schedule(groups, pnms, group_size, capacity)

    print("\nSchedule successfully generated.")
    print("Valid grouping found after {:,} attempt(s).".format(attempts))
    print("Method: {}".format(method))
    print("\nRandomization seed: {}".format(seed))

    info = {
        "seed": seed,
        "generated": datetime.now().strftime("%Y-%m-%d %I:%M %p"),
        "pnm_count": len(pnms),
        "group_size": group_size,
        "group_count": len(groups),
        "duration": duration,
        "slot_count": len(slots),
        "capacity": capacity,
        "attempts": attempts,
        "method": method,
        "csv_path": os.path.basename(csv_path),
    }
    write_excel(output_path, groups, pnms, slots, info)

    print("\nOutput saved to:")
    print(os.path.abspath(output_path))
    print("\nOpen the file and assign five brothers to each group by hand in the")
    print("Brother 1-5 columns. Keep the seed above if you need to prove the")
    print("groups were generated at random.\n")
    return 0


def main():
    try:
        return run(parse_args())
    except SchedulerError as error:
        print("\n" + "=" * 40)
        print("STOPPED - PLEASE FIX THE FOLLOWING")
        print("=" * 40 + "\n")
        print(str(error))
        print("")
        return 1
    except KeyboardInterrupt:
        print("\n\nCancelled. Nothing was saved.\n")
        return 1


if __name__ == "__main__":
    sys.exit(main())
